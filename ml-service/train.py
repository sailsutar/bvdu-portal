"""
Train Random Forest models for BVU Portal student status prediction.

Outputs:
  - models/rf_status_classifier.joblib  → At Risk | Average | Good
  - models/rf_avg_regressor.joblib      → predicted overall average %
  - models/label_encoder.joblib
  - models/feature_medians.joblib
  - models/metrics.json

Usage:
  python train.py --excel path/to/TY_marks.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.ensemble import RandomForestClassifier, RandomForestRegressor
from sklearn.metrics import (
    accuracy_score,
    classification_report,
    confusion_matrix,
    mean_absolute_error,
    r2_score,
)
from sklearn.model_selection import cross_val_score, train_test_split
from sklearn.preprocessing import LabelEncoder

PASS_UT = 8
FEATURE_COLS = [
    "attendance",
    "mean_ut1",
    "mean_ut2",
    "mean_reut",
    "mean_subject_avg",
    "mean_subject_pct",
    "n_subjects",
    "n_ut1_present",
    "n_ut2_present",
    "fail_count",
    "fail_rate",
]

ASSESS_SUFFIXES = [
    ("RE-UT-20", "RE-UT-20"),
    ("RE_UT_20", "RE-UT-20"),
    ("RE-UT", "RE-UT"),
    ("RE_UT", "RE-UT"),
    ("UT-1", "UT-1"),
    ("UT_1", "UT-1"),
    ("UT-2", "UT-2"),
    ("UT_2", "UT-2"),
    ("LAB_Pct", "LAB_Pct"),
    ("LAB_Total", "LAB_Total"),
    ("Average", "Average"),
    ("Avg", "Average"),
    ("Pct", "Pct"),
    ("Percentage", "Pct"),
    ("Total", "Total"),
]


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if np.isfinite(v) else None
    s = str(v).strip()
    if not s or re.match(r"^(defaulter|ab|absent|na|n/a|#num!|#value!|#ref!|-)$", s, re.I):
        return None
    m = re.search(r"-?\d+(\.\d+)?", s.replace(",", ""))
    return float(m.group(0)) if m else None


def parse_col(h: str):
    h = (h or "").strip()
    for suf, canon in sorted(ASSESS_SUFFIXES, key=lambda x: -len(x[0])):
        if re.search(r"[_\s]" + re.escape(suf) + r"$", h, re.I) or h.upper().endswith(suf.upper()):
            subj = re.sub(r"[_\s]*" + re.escape(suf) + r"$", "", h, flags=re.I)
            subj = subj.replace("_", " ").strip()
            if subj:
                return subj, canon
    return None, None


def label_status(attendance, avg_pct, fail_count, n_assessments):
    att = attendance if attendance is not None else 50
    pct = avg_pct if avg_pct is not None else 50
    fail_rate = fail_count / max(n_assessments, 1)
    if att < 65 or pct < 50 or fail_rate >= 0.4:
        return "At Risk"
    if att >= 75 and pct >= 72 and fail_rate <= 0.15:
        return "Good"
    return "Average"


def load_students(excel_path: str) -> list[dict]:
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_idx = 0
    for i, row in enumerate(rows[:8]):
        cells = [str(c or "").strip() for c in row]
        if any(c.upper() == "PRN" for c in cells):
            header_idx = i
            break
    headers = [str(c or "").strip() if c is not None else f"Column_{j}" for j, c in enumerate(rows[header_idx])]
    data_rows = rows[header_idx + 1 :]

    meta, academic = {}, []
    for i, h in enumerate(headers):
        hl = h.lower().replace(" ", "_")
        if hl == "prn":
            meta["prn"] = i
        elif "student" in hl and "name" in hl or hl in ("name", "student_name"):
            meta["name"] = i
        elif hl in ("section", "class", "cls"):
            meta["section"] = i
        elif "overall" in hl and "attend" in hl or hl == "overall_attendance":
            meta["attendance"] = i
        elif hl in ("average_pct", "avg_pct"):
            meta["avg_pct"] = i
        else:
            subj, assess = parse_col(h)
            if subj and assess:
                academic.append((i, subj, assess))

    students = {}
    for row in data_rows:
        if not row or meta.get("prn") is None or row[meta["prn"]] is None:
            continue
        prn = str(row[meta["prn"]]).strip()
        if not prn or prn.lower() in ("none", "prn"):
            continue
        if prn not in students:
            students[prn] = {
                "prn": prn,
                "name": str(row[meta["name"]]).strip() if meta.get("name") is not None else "",
                "attendance": None,
                "avg_pct": None,
                "marks": {},
            }
        st = students[prn]
        if meta.get("attendance") is not None:
            a = to_num(row[meta["attendance"]])
            if a is not None:
                st["attendance"] = a
        if meta.get("avg_pct") is not None:
            a = to_num(row[meta["avg_pct"]])
            if a is not None:
                st["avg_pct"] = a
        for idx, subj, assess in academic:
            if idx >= len(row):
                continue
            val = to_num(row[idx])
            if val is None:
                continue
            st["marks"].setdefault(subj, {})[assess] = val
    return list(students.values())


def build_feature_frame(students: list[dict]) -> pd.DataFrame:
    records = []
    for st in students:
        marks = st["marks"]
        ut1s, ut2s, avgs, pcts, reuts = [], [], [], [], []
        fail_count = 0
        for ass in marks.values():
            if "UT-1" in ass:
                ut1s.append(ass["UT-1"])
                if ass["UT-1"] < PASS_UT:
                    fail_count += 1
            if "UT-2" in ass:
                ut2s.append(ass["UT-2"])
                if ass["UT-2"] < PASS_UT:
                    fail_count += 1
            if "RE-UT" in ass:
                reuts.append(ass["RE-UT"])
            if "Average" in ass:
                avgs.append(ass["Average"])
            if "Pct" in ass:
                pcts.append(ass["Pct"])
        n_assess = max(len(ut1s) + len(ut2s), 1)
        avg_pct = st["avg_pct"] if st["avg_pct"] is not None else (float(np.mean(pcts)) if pcts else None)
        feat = {
            "prn": st["prn"],
            "name": st.get("name", ""),
            "attendance": st["attendance"] if st["attendance"] is not None else np.nan,
            "mean_ut1": float(np.mean(ut1s)) if ut1s else np.nan,
            "mean_ut2": float(np.mean(ut2s)) if ut2s else np.nan,
            "mean_reut": float(np.mean(reuts)) if reuts else np.nan,
            "mean_subject_avg": float(np.mean(avgs)) if avgs else np.nan,
            "mean_subject_pct": float(np.mean(pcts)) if pcts else np.nan,
            "n_subjects": len(marks),
            "n_ut1_present": len(ut1s),
            "n_ut2_present": len(ut2s),
            "fail_count": fail_count,
            "fail_rate": fail_count / n_assess,
            "avg_pct": avg_pct if avg_pct is not None else np.nan,
        }
        feat["status"] = label_status(feat["attendance"], feat["avg_pct"], fail_count, n_assess)
        records.append(feat)
    return pd.DataFrame(records)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", required=True, help="Path to academic Excel file")
    parser.add_argument("--out", default=str(Path(__file__).parent / "models"))
    args = parser.parse_args()

    students = load_students(args.excel)
    df = build_feature_frame(students)
    if df.empty:
        raise SystemExit("No students parsed from Excel")

    X = df[FEATURE_COLS].copy()
    medians = X.median()
    X = X.fillna(medians)
    y_class = df["status"]
    y_reg = df["avg_pct"].fillna(df["avg_pct"].median())

    le = LabelEncoder()
    y_enc = le.fit_transform(y_class)

    X_train, X_test, y_train, y_test, yr_train, yr_test = train_test_split(
        X, y_enc, y_reg, test_size=0.25, random_state=42, stratify=y_enc
    )

    clf = RandomForestClassifier(
        n_estimators=200, max_depth=8, min_samples_leaf=2, class_weight="balanced", random_state=42
    )
    clf.fit(X_train, y_train)
    y_pred = clf.predict(X_test)
    acc = accuracy_score(y_test, y_pred)
    cv = cross_val_score(clf, X, y_enc, cv=5, scoring="accuracy")

    reg = RandomForestRegressor(n_estimators=200, max_depth=8, min_samples_leaf=2, random_state=42)
    reg.fit(X_train, yr_train)
    yr_pred = reg.predict(X_test)
    mae = mean_absolute_error(yr_test, yr_pred)
    r2 = r2_score(yr_test, yr_pred)

    print("Classes:", list(le.classes_))
    print("Holdout accuracy:", round(acc, 4))
    print(classification_report(y_test, y_pred, target_names=le.classes_))
    print("Confusion matrix:\n", confusion_matrix(y_test, y_pred))
    print("CV accuracy:", round(cv.mean(), 4), "+/-", round(cv.std(), 4))
    print("Regressor MAE:", round(mae, 3), "R2:", round(r2, 4))

    clf.fit(X, y_enc)
    reg.fit(X, y_reg)

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    joblib.dump(clf, out / "rf_status_classifier.joblib")
    joblib.dump(reg, out / "rf_avg_regressor.joblib")
    joblib.dump(le, out / "label_encoder.joblib")
    joblib.dump(medians.to_dict(), out / "feature_medians.joblib")
    metrics = {
        "feature_cols": FEATURE_COLS,
        "classes": list(le.classes_),
        "accuracy_holdout": float(acc),
        "cv_accuracy_mean": float(cv.mean()),
        "cv_accuracy_std": float(cv.std()),
        "regressor_mae": float(mae),
        "regressor_r2": float(r2),
        "n_samples": int(len(df)),
        "status_counts": df["status"].value_counts().to_dict(),
        "pass_ut_threshold": PASS_UT,
    }
    (out / "metrics.json").write_text(json.dumps(metrics, indent=2))
    print("Saved to", out)


if __name__ == "__main__":
    main()
